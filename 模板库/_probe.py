from docx import Document
from docx.oxml.ns import qn
d = Document('试验过程点检记录.docx')
print('paras', len(d.paragraphs), 'tables', len(d.tables))
for i, p in enumerate(d.paragraphs):
    if p.text.strip():
        print('P%d:' % i, repr(p.text[:80]))
for ti, t in enumerate(d.tables):
    print('--- TABLE %d: %d rows x %d cols ---' % (ti, len(t.rows), len(t.columns)))
    for ri, row in enumerate(t.rows):
        seen = []
        txts = []
        for c in row.cells:
            if c._tc not in seen:
                seen.append(c._tc)
                txts.append(c.text.strip()[:24])
        print('  R%d:' % ri, ' | '.join(txts))
